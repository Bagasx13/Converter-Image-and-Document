import os
import sqlite3
from datetime import datetime
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.middleware.cors import CORSMiddleware
import shutil
import html
from pypdf import PdfReader
from docx import Document
import openpyxl
from pptx import Presentation
from PIL import Image
from pydantic import BaseModel

# Import ReportLab untuk pembuatan PDF
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

app = FastAPI(title="Document & Image Converter API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = "uploads"
OUTPUT_DIR = "outputs"
DB_FILE = "converter.db"

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS conversions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            original_filename TEXT NOT NULL,
            source_format TEXT NOT NULL,
            target_format TEXT NOT NULL,
            file_size INTEGER NOT NULL,
            status TEXT DEFAULT 'completed',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)
    # Tabel Pengguna (Users) untuk Autentikasi
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL
        )
    """)
    conn.commit()
    conn.close()

init_db()

class UserAuth(BaseModel):
    username: str
    password: str

@app.post("/api/register")
async def register_user(user: UserAuth):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO users (username, password) VALUES (?, ?)", (user.username, user.password))
        conn.commit()
        return {"status": "success", "message": "Registrasi berhasil! Silakan login."}
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=400, detail="Username sudah digunakan.")
    finally:
        conn.close()

@app.post("/api/login")
async def login_user(user: UserAuth):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE username = ? AND password = ?", (user.username, user.password))
    db_user = cursor.fetchone()
    conn.close()
    
    if db_user:
        return {"status": "success", "message": "Login berhasil", "username": user.username}
    raise HTTPException(status_code=401, detail="Username atau password salah.")

@app.post("/api/convert")
async def convert_document(
    file: UploadFile = File(...),
    target_format: str = Form(...)
):
    filename = file.filename
    ext = filename.split(".")[-1].lower()
    target_format = target_format.lower()
    
    input_path = os.path.join(UPLOAD_DIR, filename)
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    
    file_size = os.path.getsize(input_path)
    base_name = os.path.splitext(filename)[0]
    
    output_filename = f"converted_{base_name}.{target_format}"
    output_path = os.path.join(OUTPUT_DIR, output_filename)

    supported_image_exts = ["png", "jpg", "jpeg", "webp"]

    try:
        if ext == target_format:
            shutil.copy(input_path, output_path)
        elif ext in supported_image_exts and target_format in supported_image_exts:
            with Image.open(input_path) as img:
                if target_format in ["jpg", "jpeg"] and img.mode in ("RGBA", "P"):
                    img = img.convert("RGB")
                img.save(output_path)
        elif ext in supported_image_exts and target_format == "pdf":
            with Image.open(input_path) as img:
                rgb_img = img.convert("RGB")
                rgb_img.save(output_path, "PDF", resolution=100.0)
        elif target_format == "pdf":
            pdf_doc = SimpleDocTemplate(
                output_path, 
                pagesize=landscape(letter), 
                rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
            )
            styles = getSampleStyleSheet()
            story = []

            if ext == "docx":
                doc = Document(input_path)
                story.append(Paragraph("<b>Dokumen Word (.docx)</b>", styles['Heading1']))
                story.append(Spacer(1, 10))
                for p in doc.paragraphs:
                    if p.text.strip():
                        story.append(Paragraph(html.escape(p.text), styles['Normal']))
                        story.append(Spacer(1, 4))
            elif ext == "xlsx":
                wb = openpyxl.load_workbook(input_path, data_only=True)
                sheet = wb.active
                story.append(Paragraph("<b>Dokumen Excel (.xlsx)</b>", styles['Heading1']))
                story.append(Spacer(1, 10))
                data = []
                for row in sheet.iter_rows(values_only=True):
                    row_data = [html.escape(str(cell)) if cell is not None else "" for cell in row]
                    if any(row_data):
                        data.append(row_data[:8])
                if data:
                    t = Table(data[:60])
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('FONTSIZE', (0, 0), (-1, -1), 8)
                    ]))
                    story.append(t)
            elif ext in ["ppt", "pptx"]:
                prs = Presentation(input_path)
                for idx, slide in enumerate(prs.slides, 1):
                    story.append(Paragraph(f"<b>SLIDE {idx}</b>", styles['Heading1']))
                    story.append(Spacer(1, 8))
                    has_content = False
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    story.append(Paragraph(html.escape(p.text), styles['Normal']))
                                    story.append(Spacer(1, 4))
                                    has_content = True
                    if not has_content:
                        story.append(Paragraph("<i>(Slide kosong atau berupa gambar)</i>", styles['Normal']))
                    story.append(Spacer(1, 10))
                    story.append(PageBreak())
            elif ext == "pdf":
                reader = PdfReader(input_path)
                story.append(Paragraph("<b>Dokumen PDF Original</b>", styles['Heading1']))
                story.append(Spacer(1, 10))
                for idx, page in enumerate(reader.pages, 1):
                    extracted = page.extract_text()
                    if extracted:
                        story.append(Paragraph(f"<b>Halaman {idx}</b>", styles['Heading2']))
                        story.append(Paragraph(html.escape(extracted), styles['Normal']))
                        story.append(Spacer(1, 8))
                    story.append(PageBreak())

            if story:
                pdf_doc.build(story)
            else:
                raise Exception("Gagal memproses dokumen kosong.")
        elif target_format == "txt":
            text_content = ""
            if ext == "docx":
                doc = Document(input_path)
                text_content = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            elif ext == "xlsx":
                wb = openpyxl.load_workbook(input_path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(values_only=True):
                    text_content += "\t".join([str(c) if c is not None else "" for c in row]) + "\n"
            elif ext in ["ppt", "pptx"]:
                prs = Presentation(input_path)
                for idx, slide in enumerate(prs.slides, 1):
                    text_content += f"\n=== SLIDE {idx} ===\n"
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    text_content += p.text + "\n"
            elif ext == "pdf":
                reader = PdfReader(input_path)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text_content += extracted + "\n"
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(text_content)
        else:
            shutil.copy(input_path, output_path)

        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO conversions (original_filename, source_format, target_format, file_size, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (filename, ext, target_format, file_size, "completed", datetime.now()))
        conn.commit()
        conn.close()

        return {
            "status": "success",
            "message": f"File berhasil dikonversi ke .{target_format}",
            "download_url": f"/api/download/{output_filename}"
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/history")
async def get_conversion_history():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM conversions ORDER BY created_at DESC LIMIT 15")
    rows = cursor.fetchall()
    conn.close()
    return {"history": [dict(row) for row in rows]}

@app.delete("/api/history/{history_id}")
async def delete_history_item(history_id: int):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM conversions WHERE id = ?", (history_id,))
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Riwayat berhasil dihapus"}

@app.delete("/api/history")
async def clear_all_history():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM conversions")
    conn.commit()
    conn.close()
    return {"status": "success", "message": "Semua riwayat berhasil dibersihkan"}

@app.get("/api/download/{filename}")
async def download_file(filename: str):
    file_path = os.path.join(OUTPUT_DIR, filename)
    if os.path.exists(file_path):
        return FileResponse(file_path, media_type="application/octet-stream", filename=filename)
    raise HTTPException(status_code=404, detail="File tidak ditemukan.")

@app.get("/", response_class=HTMLResponse)
async def serve_home():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)